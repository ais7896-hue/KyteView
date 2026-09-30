import os
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

from PySide6.QtWidgets import QApplication

app = QApplication.instance() or QApplication(sys.argv)

from core.preview_router import get_renderer
from ui.preview_window import PreviewWindow
from ui.settings_dialog import SettingsDialog
from ui.license_dialog import LicenseDialog


class TestAllRenderers(unittest.TestCase):
    def setUp(self):
        self.td = tempfile.TemporaryDirectory()
        self.tdp = Path(self.td.name)

    def tearDown(self):
        try:
            self.td.cleanup()
        except Exception:
            pass

    def _render(self, path: Path):
        r = get_renderer(path)
        w = r.render(path)
        self.assertIsNotNone(w)
        if hasattr(r, "cleanup"):
            r.cleanup()
        if w is not None:
            w.deleteLater()

    def test_text_and_code(self):
        txt = self.tdp / "test.txt"
        txt.write_text("Hello KyteView\nLine 2", encoding="utf-8")
        self._render(txt)

        py_f = self.tdp / "test.py"
        py_f.write_text("def hello():\n    print('world')\n", encoding="utf-8")
        self._render(py_f)

        json_f = self.tdp / "test.json"
        json_f.write_text('{"name": "KyteView", "version": "1.0.0"}', encoding="utf-8")
        self._render(json_f)

    def test_tables(self):
        csv_f = self.tdp / "test.csv"
        csv_f.write_text("Col1,Col2,Col3\n1,2,3\n4,5,6", encoding="utf-8")
        self._render(csv_f)

        xlsx_f = self.tdp / "test.xlsx"
        import xlsxwriter
        workbook = xlsxwriter.Workbook(str(xlsx_f))
        worksheet = workbook.add_worksheet()
        worksheet.write('A1', 'Hello')
        worksheet.write('B1', 'World')
        workbook.close()
        self._render(xlsx_f)

    def test_archive(self):
        zip_f = self.tdp / "test.zip"
        with zipfile.ZipFile(zip_f, "w") as zf:
            zf.writestr("inside.txt", "inside content")
            zf.writestr("sub/inside2.txt", "nested content")
        self._render(zip_f)

    def test_media_and_docs(self):
        md_f = self.tdp / "test.md"
        md_f.write_text("# Title\n- Item 1\n- Item 2", encoding="utf-8")
        self._render(md_f)

        from PIL import Image
        png_f = self.tdp / "test.png"
        img = Image.new("RGB", (100, 100), color="blue")
        img.save(png_f)
        self._render(png_f)

        from pptx import Presentation
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text = "Hello Presentation"
        pptx_f = self.tdp / "test.pptx"
        prs.save(pptx_f)
        self._render(pptx_f)

        docx_f = self.tdp / "test.docx"
        with zipfile.ZipFile(docx_f, "w") as zf:
            zf.writestr("[Content_Types].xml", '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
            zf.writestr("_rels/.rels", '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>')
            zf.writestr("word/document.xml", '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>Hello KyteView Docx</w:t></w:r></w:p></w:body></w:document>')
        self._render(docx_f)

        import pymupdf
        pdf_doc = pymupdf.open()
        page = pdf_doc.new_page()
        page.insert_text((50, 50), "Hello PDF from KyteView")
        pdf_f = self.tdp / "test.pdf"
        pdf_doc.save(pdf_f)
        pdf_doc.close()
        self._render(pdf_f)

    def test_ui_components(self):
        pw = PreviewWindow()
        txt = self.tdp / "ui_test.txt"
        txt.write_text("UI Test", encoding="utf-8")
        pw.show_file(txt, "1 / 1", "")
        pw.hide_window()

        docx_f = self.tdp / "test.docx"
        pw.show_file(docx_f, "1 / 1", "")
        pw.hide_window()

        sd = SettingsDialog(None)
        sd._on_save_clicked()

        ld = LicenseDialog(None)
        ld.refresh_ui_state()
        ld.reject()


if __name__ == "__main__":
    unittest.main()
